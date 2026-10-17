#!/usr/bin/env python3
"""Curate the WINEP measures the importer can't auto-link (descriptive-named BW actions, works-named
monitoring, SO-vs-SSO outlet mismatches). Two phases:

  1) SUGGEST (default): list every in-scope action that is NOT yet in winep_actions, each with its
     best outlet match, and write an editable CSV.
         export DATABASE_URL='<instance URL>'
         python3 winep_curate.py --config config/catchments/dart.json "~/Downloads/PR24 WINEP National Dataset.xlsx"
     -> writes winep_curation_<river>.csv. Open it, set confirm=y on the rows you accept (and fix
        suggested_asset_id where the suggestion is wrong).

  2) EMIT: turn the confirmed CSV into SQL (insert the action with source 'winep_pr24_curated' +
     a winep_asset_links row to the chosen outlet). Apply with run_sql.py.
         python3 winep_curate.py --config config/catchments/dart.json "~/Downloads/PR24 WINEP National Dataset.xlsx" --emit winep_curation_dart.csv > dart_curated.sql
         python3 run_sql.py dart_curated.sql

Curated rows use a distinct source so the official loader never deletes them; if a later official run
can link the same action, its upsert converts the row and the winep_asset_links row survives (its id
is preserved). Read-only on the DB in suggest mode.
"""
import csv, math, os, sys

url = (os.environ.get("DATABASE_URL") or os.environ.get("DB_URL") or "").strip()
if not url:
    sys.exit("set DATABASE_URL first")

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "scripts"))
import catchment_config
from import_winep import build_record, _drivers, q, SHEET
from openpyxl import load_workbook
import psycopg2

PROX_SUGGEST_M = 2000   # how far from one of our outlets a geocoded action is still worth suggesting


def _cfg_arg():
    for i, v in enumerate(sys.argv):
        if v == "--config" and i + 1 < len(sys.argv):
            return sys.argv[i + 1]
    return None


def haversine_m(a_lat, a_lon, b_lat, b_lon):
    R = 6371000.0
    p1, p2 = math.radians(a_lat), math.radians(b_lat)
    dp, dl = math.radians(b_lat - a_lat), math.radians(b_lon - a_lon)
    h = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * R * math.asin(math.sqrt(h))


def first_seg(name):
    return (name or "").split("_")[0].strip().upper()


def load_file_records(path, company):
    wb = load_workbook(os.path.expanduser(path), read_only=True, data_only=True)
    it = wb[SHEET].iter_rows(values_only=True)
    hdr = [str(c).strip() if c is not None else "" for c in next(it)]
    idx = {h: i for i, h in enumerate(hdr)}
    ci = idx.get("Water_Company")
    dmap = _drivers()
    recs = {}
    for r in it:
        if ci is None or ci >= len(r) or company.upper() not in str(r[ci] or "").upper():
            continue
        rec = build_record(r, idx, dmap)
        if rec:
            recs[(rec["action_id"], rec["comp"])] = rec
    return recs


def connect():
    conn = psycopg2.connect(url)
    cur = conn.cursor()
    cur.execute("select public_org from app_config limit 1")
    return conn, cur, cur.fetchone()[0]


def suggest(cfg, company, csv_path):
    conn, cur, org = connect()
    cur.execute("""select id, asset_name, latitude, longitude from sewage_assets
                   where organisation_id=%s and asset_name is not null""", (org,))
    assets = [{"id": r[0], "name": r[1], "lat": r[2], "lon": r[3]} for r in cur.fetchall()]
    by_seg = {}
    for a in assets:
        by_seg.setdefault(first_seg(a["name"]), []).append(a)
    cur.execute("""select action_id, coalesce(action_component,'') from winep_actions
                   where organisation_id=%s and cycle='PR24'""", (org,))
    imported = {(a, c) for a, c in cur.fetchall()}
    conn.close()

    recs = load_file_records(_xlsx(), company)
    out = []
    for (aid, comp), rec in recs.items():
        if (aid, comp) in imported:
            continue
        seg = first_seg(rec["name"])
        works = by_seg.get(seg, []) if seg else []
        lat, lon = rec["lat"], rec["lon"]
        near = None
        if lat is not None:
            cand = sorted(((haversine_m(lat, lon, a["lat"], a["lon"]), a)
                           for a in assets if a["lat"] is not None), key=lambda x: x[0])
            near = cand[0] if cand else None
        if works:
            # prefer the works asset nearest the action (if geocoded), else the first
            if lat is not None:
                works = sorted(works, key=lambda a: (haversine_m(lat, lon, a["lat"], a["lon"])
                                                     if a["lat"] is not None else 9e9))
            pick, match, dist = works[0], "works-seg", (
                round(haversine_m(lat, lon, works[0]["lat"], works[0]["lon"]))
                if lat is not None and works[0]["lat"] is not None else "")
            alt = ";".join(a["name"] for a in works[1:4])
        elif near and near[0] <= PROX_SUGGEST_M:
            pick, match, dist = near[1], f"prox", round(near[0])
            alt = ""
        else:
            continue  # not plausibly ours — skip (neighbouring-catchment noise)
        out.append({
            "confirm": "", "action_id": aid, "action_component": comp,
            "driver_code": rec["dcode"] or "", "due": rec["due"] or "",
            "action_name": (rec["name"] or "").replace("\n", " "),
            "match": match, "distance_m": dist,
            "suggested_asset_id": pick["id"], "suggested_asset_name": pick["name"],
            "alt_assets": alt,
        })

    out.sort(key=lambda x: (0 if x["match"] == "works-seg" else 1, str(x["distance_m"])))
    cols = ["confirm", "action_id", "action_component", "driver_code", "due", "action_name",
            "match", "distance_m", "suggested_asset_id", "suggested_asset_name", "alt_assets"]
    with open(csv_path, "w", newline="", encoding="utf-8") as fh:
        wr = csv.DictWriter(fh, fieldnames=cols)
        wr.writeheader()
        wr.writerows(out)
    print(f"{len(out)} action(s) to review -> {csv_path}", file=sys.stderr)
    print(f"{'match':<10}{'driver':<13}{'due':<12}{'dist':>6}  action -> suggested outlet")
    for r in out:
        print(f"{r['match']:<10}{r['driver_code']:<13}{str(r['due']):<12}{str(r['distance_m']):>6}  "
              f"{r['action_name'][:44]:<44} -> {r['suggested_asset_name']}")
    print(f"\nEdit {csv_path}: set confirm=y on rows to keep (fix suggested_asset_id where wrong), "
          f"then re-run with --emit {csv_path}", file=sys.stderr)


def _flt(v):
    try:
        return repr(float(v))
    except (TypeError, ValueError):
        return "null"


def emit(cfg, company, csv_path):
    recs = load_file_records(_xlsx(), company)
    conn, cur, org = connect()
    conn.close()
    orgl = q(org) + "::uuid"
    cols = ["water_company", "driver_code", "driver_label", "driver_obligation",
            "driver_code_secondary", "driver_code_tertiary", "action_name", "action_description",
            "tier1_outcome", "options_outcome", "aim", "spatial_scale", "ea_water_body_id", "wb_type",
            "wb_name", "completion_date", "bathing_water", "shellfish_water", "sssi", "sac_spa_ramsar",
            "mcz", "current_permit_dwf", "proposed_permit_dwf", "current_bod", "proposed_bod",
            "current_nh3", "proposed_nh3", "current_p", "proposed_p", "latitude", "longitude"]
    keymap = {  # winep_actions column -> rec key
        "water_company": "company", "driver_code": "dcode", "driver_label": "dlabel",
        "driver_obligation": "doblig", "driver_code_secondary": "dsec", "driver_code_tertiary": "dter",
        "action_name": "name", "action_description": "desc", "tier1_outcome": "tier1",
        "options_outcome": "opt", "aim": "aim", "spatial_scale": "scale", "ea_water_body_id": "ea_wb",
        "wb_type": "wb_type", "wb_name": "wb_name", "completion_date": "due", "bathing_water": "bw",
        "shellfish_water": "sw", "sssi": "sssi", "sac_spa_ramsar": "sac", "mcz": "mcz",
        "current_permit_dwf": "cdwf", "proposed_permit_dwf": "pdwf", "current_bod": "cbod",
        "proposed_bod": "pbod", "current_nh3": "cnh3", "proposed_nh3": "pnh3", "current_p": "cp",
        "proposed_p": "pp"}
    lines = ["-- Curated WINEP links (confirmed by hand). Idempotent: action upsert is a no-op if the",
             "-- row already exists; links are on-conflict-do-nothing. Source 'winep_pr24_curated' so",
             "-- the official loader never deletes them.", "begin;"]
    n = 0
    with open(csv_path, encoding="utf-8") as fh:
        for row in csv.DictReader(fh):
            if str(row.get("confirm", "")).strip().lower() not in ("y", "yes", "1", "true"):
                continue
            key = (row["action_id"].strip(), row["action_component"].strip())
            rec = recs.get(key)
            if not rec:
                sys.stderr.write(f"-- skip {key}: not found in file\n")
                continue
            asset = row["suggested_asset_id"].strip()
            if not asset:
                sys.stderr.write(f"-- skip {key}: no suggested_asset_id\n")
                continue
            vals = [orgl, "'PR24'", q(rec["action_id"]), q(rec["comp"])]
            for c in cols:
                if c in ("latitude", "longitude"):
                    vals.append(_flt(rec["lon" if c == "longitude" else "lat"]))
                else:
                    vals.append(q(rec[keymap[c]]))
            vals.append("'winep_pr24_curated'")
            collist = "organisation_id, cycle, action_id, action_component, " + ", ".join(cols) + ", source"
            lines.append(f"insert into winep_actions ({collist})\n  values ({', '.join(vals)})\n"
                         f"  on conflict (organisation_id, cycle, action_id, action_component) do nothing;")
            lines.append(f"""insert into winep_asset_links (organisation_id, winep_action_id, asset_id, note)
  select {orgl}, wa.id, {q(asset)}::uuid, 'curated: {row.get('match','manual')}'
  from winep_actions wa
  where wa.organisation_id={orgl} and wa.cycle='PR24'
    and wa.action_id={q(rec['action_id'])} and wa.action_component={q(rec['comp'])}
  on conflict (winep_action_id, asset_id) do nothing;""")
            n += 1
    lines.append(f"select 'curated links written: ' || {n};")
    lines.append("commit;")
    print("\n".join(lines))
    print(f"-- {n} confirmed action(s) emitted", file=sys.stderr)


def _xlsx():
    args = [a for a in sys.argv[1:] if not a.startswith("--") and a != _cfg_arg()
            and a != (sys.argv[sys.argv.index("--emit") + 1] if "--emit" in sys.argv else None)]
    if not args:
        sys.exit("give the path to the official PR24 WINEP National Dataset .xlsx")
    return args[0]


def main():
    cfg = catchment_config.load()
    company = cfg["company"].get("annual_sheet_match", cfg["company"]["name"])
    if "--emit" in sys.argv:
        emit(cfg, company, sys.argv[sys.argv.index("--emit") + 1])
    else:
        csv_path = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                f"winep_curation_{cfg['river'].lower().replace(' ', '_')}.csv")
        suggest(cfg, company, csv_path)


if __name__ == "__main__":
    main()
